from pathlib import Path
from docx import Document

root = Path(__file__).resolve().parent.parent
file_path = root / 'docs' / 'Modelo_de_Dados_NoSQL_MongoDB.docx'

if file_path.exists():
    doc = Document(str(file_path))
else:
    doc = Document()

paragraphs = [
    'Modelo de Dados NoSQL MongoDB — atualização de implementação',
    '',
    'Coleções principais: users, modulos, configuracao, notas_saida',
    '',
    'users: _id, nome, username, perfil, ativo, password_hash, assinatura, data_criacao',
    'modulos: _id, codigo, nome, ordem, ativo',
    'configuracao: _id, chave, valor, updated_at',
    'notas_saida: _id, numero_referencia, funcionario, email_funcionario, departamento, motivo, estado, criado_por, aprovado_por, datas, itens, assinaturas, historico, comentario_decisao, pdf_path',
    '',
    'Estado da nota: rascunho, pendente_aprovacao, em_revisao, rejeitada, concluida',
    'Histórico: {utilizador, acao, data_hora}',
    'Assinaturas: {entregue, recebido, seguranca, aprovador}',
    'Observação: esta estrutura foi validada em testes reais contra o MongoDB Docker do projeto.',
]

for text in paragraphs:
    if text == '':
        doc.add_paragraph('')
    else:
        doc.add_paragraph(text)

doc.save(str(file_path))
print(f'DOCX_UPDATED: {file_path}')
